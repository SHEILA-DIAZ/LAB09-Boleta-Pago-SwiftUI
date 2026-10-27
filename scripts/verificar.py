"""Comprueba resultados oficiales, fórmulas del XLSX y equivalencia con Swift/Excel."""
from pathlib import Path
from decimal import Decimal as D
import hashlib
import json
import sys
import unittest
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / '.tools'), str(ROOT / 'scripts')]
import openpyxl
from generar import calcular, LABELS

FIXTURES = json.loads((ROOT/'Tests/fixtures.json').read_text(encoding='utf-8'))
LIBRO = ROOT/'Excel/Boleta_Pago_Semana09.xlsx'
OFFICIAL = [
    ['2500','0','0','0','2500','325','0','0','0','0','325','2175','225','2725'],
    ['1620','123','0','0','1743','226.59','0','0','0','0','226.59','1516.41','156.87','1899.87'],
    ['3000','123','62.50','67.50','3253','0','325.30','52.05','43.92','200','621.27','2631.73','292.77','3545.77'],
]

class Verificacion(unittest.TestCase):
    def test_oficiales(self):
        for i, expected in enumerate(OFFICIAL):
            # Ceros y totales derivados del formato; conceptos publicados contrastados al centavo.
            with self.subTest(caso=i+1): self.assertEqual(calcular(FIXTURES[i]['input']), list(map(D,expected)))

    def test_caso4(self):
        esperado = ['3360','123','43.75','47.25','3574','0','357.40','57.18','48.25','100','562.83','3011.17','321.66','3895.66']
        self.assertEqual(calcular(FIXTURES[3]['input']),list(map(D,esperado)))

    def test_docx_preservado(self):
        doc = json.loads((ROOT/'evidencias/documento.json').read_text(encoding='utf-8'))
        self.assertEqual(doc['tablas'],9)
        self.assertEqual(hashlib.sha256((ROOT/'referencias/Actividad_Grupo_B.docx').read_bytes()).hexdigest(),doc['sha256'])
        self.assertIn('9. Datos reales y fuentes',(ROOT/'referencias/guia_extraida.md').read_text(encoding='utf-8'))
        self.assertIn('10. Entregables',(ROOT/'referencias/guia_extraida.md').read_text(encoding='utf-8'))

    def test_estructura_y_formulas(self):
        w = openpyxl.load_workbook(LIBRO, data_only=False)
        self.assertEqual(w.sheetnames,['Parámetros','Boleta','Casos','Explicación'])
        for row in range(17,31): self.assertEqual(w['Boleta'][f'B{row}'].data_type,'f')
        for row in range(4,4+len(FIXTURES)):
            for col in range(10,24): self.assertEqual(w['Casos'].cell(row,col).data_type,'f')
        self.assertGreaterEqual(len(w['Boleta'].data_validations.dataValidation),8)
        self.assertGreaterEqual(len(w.defined_names),9)
        for s in w:
            for row in s:
                for cell in row:
                    if cell.data_type=='f':
                        self.assertNotIn('#REF!',cell.value)
                        self.assertNotIn('[1]',cell.value)
        w.close()

    def test_cache_excel(self):
        # Solo verifica cache almacenada, NO demuestra recálculo nativo.
        w = openpyxl.load_workbook(LIBRO,data_only=True)
        for row,f in enumerate(FIXTURES,4):
            for col,expected in enumerate(f['expected'],10):
                self.assertEqual(D(str(w['Casos'].cell(row,col).value)).quantize(D('.01')),D(expected))
        for s in w:
            for row in s:
                for cell in row: self.assertNotEqual(cell.data_type,'e',f'Error de Excel {s.title}!{cell.coordinate}')
        w.close()

    def test_swift(self):
        path=ROOT/'evidencias/swift_resultados.json'
        if not path.exists(): self.skipTest('Swift no ejecutado; no atribuir equivalencia')
        resultados=json.loads(path.read_text(encoding='utf-8'))
        self.assertGreater(resultados['checks'],200)
        self.assertEqual(len(resultados['cases']),len(FIXTURES))
        for f,actual in zip(FIXTURES,resultados['cases']):
            self.assertEqual(actual['name'],f['name'])
            self.assertEqual(list(map(D,actual['values'])),list(map(D,f['expected'])))

    def test_excel_nativo(self):
        path=ROOT/'evidencias/excel_recalculado.json'
        if not path.exists(): self.skipTest('Excel COM no ejecutado; solo cache verificada')
        resultados=json.loads(path.read_text(encoding='utf-8-sig'))
        self.assertEqual(resultados['engine'],'Microsoft Excel COM')
        self.assertTrue(resultados['invalidDaysBlocked'])
        for grupo in ['cases','form']:
            self.assertEqual(len(resultados[grupo]),len(FIXTURES))
            for f,actual in zip(FIXTURES,resultados[grupo]):
                self.assertEqual(actual['name'],f['name'])
                self.assertEqual(list(map(D,actual['values'])),list(map(D,f['expected'])))

    def test_formulas_recalculadas(self):
        path=ROOT/'evidencias/formulas_recalculadas.json'
        if not path.exists(): self.skipTest('Motor independiente no ejecutado')
        resultados=json.loads(path.read_text(encoding='utf-8'))
        self.assertEqual(resultados['engine'],'formulas')
        self.assertEqual(resultados['conceptComparisons'],448)
        self.assertEqual(resultados['invalidInputsBlocked'],12)
        self.assertTrue(resultados['parameterMutationPassed'])
        for grupo in ['cases','form']:
            self.assertEqual(len(resultados[grupo]),len(FIXTURES))
            for f,actual in zip(FIXTURES,resultados[grupo]):
                self.assertEqual(actual['name'],f['name'])
                self.assertEqual(list(map(D,actual['values'])),list(map(D,f['expected'])))

    def test_xcode_referencias(self):
        p=(ROOT/'BoletaPago.xcodeproj/project.pbxproj').read_text(encoding='utf-8')
        for f in ['BoletaPagoApp.swift','ContentView.swift','Model/Boleta.swift']:
            self.assertTrue((ROOT/'App'/f).exists())
            self.assertIn(f,p)
        ET.parse(ROOT/'BoletaPago.xcodeproj/xcshareddata/xcschemes/BoletaPago.xcscheme')
        self.assertIn('IPHONEOS_DEPLOYMENT_TARGET = 17.0',p)

if __name__=='__main__':
    suite=unittest.defaultTestLoader.loadTestsFromTestCase(Verificacion)
    with (ROOT/'evidencias/pruebas_python.txt').open('w',encoding='utf-8') as salida:
        resultado=unittest.TextTestRunner(stream=salida,verbosity=2).run(suite)
    print((ROOT/'evidencias/pruebas_python.txt').read_text(encoding='utf-8'))
    sys.exit(0 if resultado.wasSuccessful() else 1)
